import os
import PyPDF2
import docx


def extract_text(filepath: str) -> str:
    """
    Module 1: Detects file type and extracts raw text from PDF or DOCX resumes.
    Returns extracted text as a string, or empty string on failure.
    """
    ext = os.path.splitext(filepath)[1].lower()

    if ext == '.pdf':
        return _extract_from_pdf(filepath)
    elif ext == '.docx':
        return _extract_from_docx(filepath)
    else:
        return ''


def _extract_from_pdf(filepath: str) -> str:
    """Extracts text from a PDF file using PyPDF2."""
    text_parts = []
    try:
        with open(filepath, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
    except Exception as e:
        print(f"[file_parser] PDF extraction error: {e}")
        return ''
    return '\n'.join(text_parts)


def _extract_from_docx(filepath: str) -> str:
    """Extracts text from a DOCX file using python-docx."""
    text_parts = []
    try:
        doc = docx.Document(filepath)
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)
    except Exception as e:
        print(f"[file_parser] DOCX extraction error: {e}")
        return ''
    return '\n'.join(text_parts)
